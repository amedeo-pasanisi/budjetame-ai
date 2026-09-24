"""Backup pipeline (issue #112): build a multi-sheet backup workbook carrying
the Account's whole data state.

The builder is pure (no database access) — it formats already-resolved data
whose names were resolved against the Account by the route, exactly like
`services.exports` and `services.imports`. The file is a multi-sheet workbook:

- Transactions (all types, including Opening Balances — unlike the ledger Export)
- Wallets
- Categories
- Recurring Costs
- Recurring Incomes
- Skips
- Metadata (origin marker)

Entities are keyed by name; original database ids are carried alongside for
reference (ADR-0030).
"""

from dataclasses import dataclass
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook

TRANSACTIONS_HEADER = [
    "date",
    "type",
    "amount",
    "wallet",
    "source wallet",
    "destination wallet",
    "category",
    "description",
    "location",
    "recurring cost",
    "recurring income",
    "occurrence date",
    "id",
]

WALLETS_HEADER = [
    "name",
    "type",
    "frozen",
    "id",
]

CATEGORIES_HEADER = [
    "name",
    "type",
    "icon",
    "color",
    "id",
]

RECURRING_DEFINITION_HEADER = [
    "name",
    "amount",
    "interval value",
    "interval unit",
    "start date",
    "frozen",
    "freeze date",
    "id",
]

SKIPS_HEADER = [
    "definition name",
    "definition type",
    "occurrence date",
    "id",
]

METADATA_HEADER = [
    "key",
    "value",
]


@dataclass(frozen=True, slots=True)
class BackupTransaction:
    """One row of the Transactions sheet, with names (not ids) for Wallet,
    Category, and Recurring links. All Transaction types are included,
    including Opening Balances."""

    date: str
    type: str
    amount: Decimal
    wallet: str | None
    source_wallet: str | None
    destination_wallet: str | None
    category: str | None
    description: str | None
    latitude: Decimal | None
    longitude: Decimal | None
    recurring_cost: str | None
    recurring_income: str | None
    occurrence_date: str | None
    id: int


@dataclass(frozen=True, slots=True)
class BackupWallet:
    name: str
    type: str
    frozen: bool
    id: int


@dataclass(frozen=True, slots=True)
class BackupCategory:
    name: str
    type: str
    icon: str | None
    color: str
    id: int


@dataclass(frozen=True, slots=True)
class BackupRecurringDefinition:
    name: str
    amount: Decimal
    interval_value: int
    interval_unit: str
    start_date: str
    frozen: bool
    freeze_date: str | None
    id: int


@dataclass(frozen=True, slots=True)
class BackupSkip:
    definition_name: str
    definition_type: str  # "cost" or "income"
    occurrence_date: str
    id: int


def _coord(value: Decimal | None) -> str:
    """Shortest-decimal coordinate form — same as the Export format."""
    if value is None:
        return ""
    return f"{value:.6f}".rstrip("0").rstrip(".")


def _str_or_blank(value: str | None) -> str:
    return "" if value is None else value


def _str_or_blank_date(value: str | None) -> str:
    return "" if value is None else value


def build_backup_workbook(
    transactions: list[BackupTransaction],
    wallets: list[BackupWallet],
    categories: list[BackupCategory],
    recurring_costs: list[BackupRecurringDefinition],
    recurring_incomes: list[BackupRecurringDefinition],
    skips: list[BackupSkip],
    origin: str,
) -> bytes:
    """Build the multi-sheet backup workbook.

    Each sheet has a header row followed by data rows. The Metadata sheet
    carries the origin marker (the Account's email). Entities are keyed by
    name; database ids are included for reference.
    """
    workbook = Workbook()

    # --- Metadata sheet (always first, for quick inspection) ---
    metadata_sheet = workbook.active
    assert metadata_sheet is not None  # a new Workbook always has one active sheet
    metadata_sheet.title = "Metadata"
    metadata_sheet.append(METADATA_HEADER)
    metadata_sheet.append(["origin", origin])
    metadata_sheet.append(["exported_at", __import__("datetime").datetime.now(
        __import__("datetime").timezone.utc
    ).isoformat()])

    # --- Transactions sheet ---
    tx_sheet = workbook.create_sheet("Transactions")
    tx_sheet.append(TRANSACTIONS_HEADER)
    for tx in transactions:
        location = (
            f"{_coord(tx.latitude)},{_coord(tx.longitude)}"
            if tx.latitude is not None and tx.longitude is not None
            else ""
        )
        tx_sheet.append([
            tx.date,
            tx.type,
            f"{tx.amount:.2f}",
            _str_or_blank(tx.wallet),
            _str_or_blank(tx.source_wallet),
            _str_or_blank(tx.destination_wallet),
            _str_or_blank(tx.category),
            _str_or_blank(tx.description),
            location,
            _str_or_blank(tx.recurring_cost),
            _str_or_blank(tx.recurring_income),
            _str_or_blank_date(tx.occurrence_date),
            tx.id,
        ])

    # --- Wallets sheet ---
    wallet_sheet = workbook.create_sheet("Wallets")
    wallet_sheet.append(WALLETS_HEADER)
    for w in wallets:
        wallet_sheet.append([
            w.name,
            w.type,
            str(w.frozen),
            w.id,
        ])

    # --- Categories sheet ---
    cat_sheet = workbook.create_sheet("Categories")
    cat_sheet.append(CATEGORIES_HEADER)
    for c in categories:
        cat_sheet.append([
            c.name,
            c.type,
            _str_or_blank(c.icon),
            c.color,
            c.id,
        ])

    # --- Recurring Costs sheet ---
    rc_sheet = workbook.create_sheet("Recurring Costs")
    rc_sheet.append(RECURRING_DEFINITION_HEADER)
    for rc in recurring_costs:
        rc_sheet.append([
            rc.name,
            f"{rc.amount:.2f}",
            rc.interval_value,
            rc.interval_unit,
            rc.start_date,
            str(rc.frozen),
            _str_or_blank_date(rc.freeze_date),
            rc.id,
        ])

    # --- Recurring Incomes sheet ---
    ri_sheet = workbook.create_sheet("Recurring Incomes")
    ri_sheet.append(RECURRING_DEFINITION_HEADER)
    for ri in recurring_incomes:
        ri_sheet.append([
            ri.name,
            f"{ri.amount:.2f}",
            ri.interval_value,
            ri.interval_unit,
            ri.start_date,
            str(ri.frozen),
            _str_or_blank_date(ri.freeze_date),
            ri.id,
        ])

    # --- Skips sheet ---
    skips_sheet = workbook.create_sheet("Skips")
    skips_sheet.append(SKIPS_HEADER)
    for s in skips:
        skips_sheet.append([
            s.definition_name,
            s.definition_type,
            s.occurrence_date,
            s.id,
        ])

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()